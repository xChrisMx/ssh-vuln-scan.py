#!/usr/bin/env python3
"""
ssh_vuln_scan.py - subnet-wide SSH weakness sweep, single-file tool.

Runs nmap against port 22 across a subnet, scores every host with SSH open,
and writes both insecure_ssh_hosts.csv (wide format, one row per host) and
insecure_ssh_hosts.xlsx (Overview + Report sheets) - no separate generator
script needed.

Usage:
    python ssh_vuln_scan.py [subnet] [--keep-temp] [--no-xlsx]
                             [--csv-out FILE] [--xlsx-out FILE]
    python ssh_vuln_scan.py --from-csv insecure_ssh_hosts.csv --xlsx-out out.xlsx

    subnet        defaults to 192.168.0.1/24
    --keep-temp   keep the raw nmap XML instead of deleting it
    --no-xlsx     stop after the CSV; skip the openpyxl step entirely
    --from-csv    skip scanning; rebuild the .xlsx from an existing CSV
                  (e.g. one you hand-edited) instead

Requires: nmap on PATH, Python 3.8+, openpyxl (only if generating .xlsx)
"""

import argparse
import csv
import os
import re
import subprocess
import sys
from datetime import datetime
from xml.etree import ElementTree as ET

# ---------------------------------------------------------------------------
# Weak-algorithm patterns (same definitions the earlier bash/awk version
# used - kept identical so a re-scan of the same subnet reproduces the same
# counts).
# ---------------------------------------------------------------------------
WEAK_KEX_RE = re.compile(r"diffie-hellman-group1-sha1|diffie-hellman-group-exchange-sha1")
WEAK_MAC_RE = re.compile(r"hmac-sha1|hmac-md5")
WEAK_HOSTKEYALGO_RE = re.compile(r"ssh-dss|ecdsa-sha2-nistp")
WEAK_CIPHER_RE = re.compile(r"arcfour|-cbc|3des|none")

NMAP_SCRIPTS = "sshv1,ssh2-enum-algos,ssh-hostkey,ssh-auth-methods"

REPORT_HEADERS = [
    "Scan Date", "IP Address", "Hostname", "Banner", "SSHv1 Supported",
    "Weak KEX Count", "Weak MAC Count", "Weak HostKeyAlgo Count",
    "Weak Cipher Count", "Password Auth Enabled", "Weak Algo Total",
    "Weakness Score", "Host Key Info",
]
CSV_HEADERS = [
    "ScanDate", "IP", "Hostname", "Banner", "SSHv1Supported",
    "WeakKexCount", "WeakMacCount", "WeakHostKeyAlgoCount",
    "WeakCipherCount", "PasswordAuthEnabled", "HostKeyInfo",
]


# ===========================================================================
# Scanning
# ===========================================================================

def run_nmap(subnet, xml_path):
    cmd = ["nmap", "-p", "22", "-sV", "--script", NMAP_SCRIPTS, "-oX", xml_path, subnet]
    print(f"[*] Scanning subnet {subnet} for SSH vulnerabilities...")
    print("[*] (run with sudo for a SYN scan; a plain connect scan works without it, just slower)")
    print(f"[*] Command: {' '.join(cmd)}")
    try:
        result = subprocess.run(cmd)
    except FileNotFoundError:
        print("[!] nmap not found on PATH.", file=sys.stderr)
        sys.exit(1)
    if result.returncode != 0:
        print(f"[!] nmap exited with code {result.returncode}", file=sys.stderr)
        sys.exit(result.returncode)


def _table_values(script_elem, key):
    """Structured NSE output: <table key="..."><elem>value</elem>...</table>.
    Returns None (not an empty list) if no table with this key exists at
    all, so the caller can tell "no weak algorithms" apart from "nmap didn't
    emit structured output this time"."""
    for table in script_elem.findall("table"):
        if table.get("key") == key:
            return [e.text or "" for e in table.findall("elem")]
    return None


def _parse_algos_block_fallback(output_text, header):
    """Fallback if structured <table> children are absent: scan the
    flattened `output` attribute text for the labeled block and take every
    line until the next blank line - the same block-range logic the earlier
    awk version used, ported to Python."""
    values = []
    in_block = False
    for line in output_text.splitlines():
        if header in line:
            in_block = True
            continue
        if in_block:
            if line.strip() == "":
                break
            values.append(line.strip())
    return values


def parse_ssh2_enum_algos(script_elem):
    """Returns {category: [algorithm names]} for the four categories we
    score. Tries structured <table> children first; falls back to text
    parsing only if none of the four expected tables were found at all."""
    keys = ("kex_algorithms", "server_host_key_algorithms",
            "encryption_algorithms", "mac_algorithms")
    result = {k: _table_values(script_elem, k) for k in keys}

    if all(v is None for v in result.values()):
        output = script_elem.get("output", "") or ""
        result = {k: _parse_algos_block_fallback(output, k + ":") for k in keys}

    return {k: (v or []) for k, v in result.items()}


def parse_host(host_elem, scan_date):
    """Returns a row dict, or None if this host isn't a countable SSH host
    (down, port 22 not open, or port 22 open but not actually speaking SSH)."""
    status = host_elem.find("status")
    if status is None or status.get("state") != "up":
        return None

    ip = None
    for addr in host_elem.findall("address"):
        if addr.get("addrtype") in ("ipv4", "ipv6"):
            ip = addr.get("addr")
            break
    if ip is None:
        return None

    hostname = ""
    hostnames_el = host_elem.find("hostnames")
    if hostnames_el is not None:
        # Prefer a reverse-DNS (PTR) name; fall back to any other name nmap
        # attached (e.g. one supplied on the command line) if no PTR exists.
        chosen = hostnames_el.find("hostname[@type='PTR']")
        if chosen is None:
            chosen = hostnames_el.find("hostname")
        if chosen is not None:
            hostname = chosen.get("name", "")

    port_el = None
    ports_el = host_elem.find("ports")
    if ports_el is not None:
        for p in ports_el.findall("port"):
            if p.get("portid") == "22":
                port_el = p
                break
    if port_el is None:
        return None

    state_el = port_el.find("state")
    if state_el is None or state_el.get("state") != "open":
        return None  # closed / filtered / open|filtered - nothing to audit

    service_el = port_el.find("service")
    service_name = service_el.get("name", "") if service_el is not None else ""

    banner = ""
    if service_el is not None:
        product = service_el.get("product", "")
        version = service_el.get("version", "")
        extrainfo = service_el.get("extrainfo", "")
        banner = " ".join(p for p in (product, version) if p)
        if extrainfo:
            banner = f"{banner} ({extrainfo})" if banner else f"({extrainfo})"

    sshv1 = False
    kex_weak = mac_weak = hka_weak = cipher_weak = 0
    pwauth = False
    hostkey_info = ""
    ssh_script_seen = False

    for script in port_el.findall("script"):
        sid = script.get("id")
        output = script.get("output", "") or ""

        if sid == "sshv1":
            # nmap only emits a <script> element for a vulnerability-style
            # check when it actually fired - presence IS the signal.
            sshv1 = True
            ssh_script_seen = True
        elif sid == "ssh2-enum-algos":
            algos = parse_ssh2_enum_algos(script)
            kex_weak = sum(1 for a in algos["kex_algorithms"] if WEAK_KEX_RE.search(a))
            mac_weak = sum(1 for a in algos["mac_algorithms"] if WEAK_MAC_RE.search(a))
            hka_weak = sum(1 for a in algos["server_host_key_algorithms"] if WEAK_HOSTKEYALGO_RE.search(a))
            cipher_weak = sum(1 for a in algos["encryption_algorithms"] if WEAK_CIPHER_RE.search(a))
            ssh_script_seen = True
        elif sid == "ssh-auth-methods":
            if "password" in output.lower():
                pwauth = True
            ssh_script_seen = True
        elif sid == "ssh-hostkey":
            hostkey_info = " | ".join(line.strip() for line in output.splitlines() if line.strip())
            ssh_script_seen = True

    # Guard against counting a non-SSH service that happens to be sitting on
    # port 22 as a "clean SSH host": require -sV to have named it "ssh", OR
    # at least one SSH-specific NSE script to have actually produced output
    # (i.e. something on that port answered SSH's own protocol probes).
    # Without this, a host running e.g. a relocated FTP/HTTP service on 22
    # would show zero findings across the board and inflate the "Fully
    # Hardened Hosts" KPI with a host that was never running SSH at all.
    if service_name != "ssh" and not ssh_script_seen:
        return None

    weak_total = kex_weak + mac_weak + hka_weak + cipher_weak
    score = (100 if sshv1 else 0) + weak_total

    return {
        "scan_date": scan_date, "ip": ip, "hostname": hostname, "banner": banner,
        "sshv1": sshv1, "kex": kex_weak, "mac": mac_weak, "hka": hka_weak,
        "cip": cipher_weak, "pwauth": pwauth, "hostkey_info": hostkey_info,
        "weak_total": weak_total, "score": score,
    }


def parse_xml_file(xml_path, scan_date):
    tree = ET.parse(xml_path)
    root = tree.getroot()
    rows = []
    for host_elem in root.findall("host"):
        row = parse_host(host_elem, scan_date)
        if row is not None:
            rows.append(row)
    rows.sort(key=lambda r: tuple(int(x) for x in r["ip"].split(".")) if r["ip"].count(".") == 3 else (r["ip"],))
    return rows


# ===========================================================================
# CSV
# ===========================================================================

def write_csv(rows, csv_path):
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(CSV_HEADERS)
        for r in rows:
            w.writerow([
                r["scan_date"], r["ip"], r["hostname"], r["banner"],
                "TRUE" if r["sshv1"] else "FALSE",
                r["kex"], r["mac"], r["hka"], r["cip"],
                "TRUE" if r["pwauth"] else "FALSE",
                r["hostkey_info"],
            ])
    print(f"[*] CSV saved to: {csv_path}")


def read_rows_from_csv(csv_path):
    """Rebuild the same row-dict shape parse_host() produces, from a
    previously-written (or hand-edited) CSV. Used by --from-csv."""
    rows = []
    with open(csv_path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            sshv1 = r["SSHv1Supported"].strip().upper() == "TRUE"
            pwauth = r["PasswordAuthEnabled"].strip().upper() == "TRUE"
            kex = int(r["WeakKexCount"])
            mac = int(r["WeakMacCount"])
            hka = int(r["WeakHostKeyAlgoCount"])
            cip = int(r["WeakCipherCount"])
            weak_total = kex + mac + hka + cip
            score = (100 if sshv1 else 0) + weak_total
            rows.append({
                "scan_date": r["ScanDate"], "ip": r["IP"], "hostname": r["Hostname"],
                "banner": r["Banner"], "sshv1": sshv1, "kex": kex, "mac": mac,
                "hka": hka, "cip": cip, "pwauth": pwauth,
                "hostkey_info": r.get("HostKeyInfo", ""),
                "weak_total": weak_total, "score": score,
            })
    return rows


# ===========================================================================
# XLSX report
# ===========================================================================

def build_workbook(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.marker import DataPoint
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.utils import get_column_letter

    RED_FILL = PatternFill("solid", fgColor="FFC7CE")
    RED_FONT = Font(color="9C0006", bold=True)
    AMBER_FILL = PatternFill("solid", fgColor="FFEB9C")
    AMBER_FONT = Font(color="9C6500")
    GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
    GREEN_FONT = Font(color="006100")
    NEUTRAL_FILL = PatternFill("solid", fgColor="DCE6F1")
    HEADER_FILL = PatternFill("solid", fgColor="1F3864")
    HEADER_FONT = Font(color="FFFFFF", bold=True, size=18)
    SECTION_FONT = Font(bold=True, size=13)
    LABEL_FONT = Font(bold=True, size=10, color="595959")
    SUBTEXT_FONT = Font(italic=True, size=9, color="808080")
    TILE_NUM_FONT = Font(bold=True, size=32)
    THIN_BORDER = Border(bottom=Side(style="thin", color="BFBFBF"))
    RED_BAR = "FFC7CE"
    AMBER_BAR = "FFEB9C"

    COL = {name: i + 1 for i, name in enumerate(REPORT_HEADERS)}

    def build_report_sheet(wb):
        ws = wb.create_sheet("Report")
        ws["A1"] = "SSH Weakness Scan - Detailed Report"
        ws["A1"].font = Font(bold=True, size=14)

        for c, name in enumerate(REPORT_HEADERS, start=1):
            cell = ws.cell(row=2, column=c, value=name)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9D9D9")

        for i, r in enumerate(rows):
            row_num = 3 + i
            try:
                date_val = datetime.strptime(r["scan_date"], "%Y-%m-%d").date()
            except (ValueError, TypeError):
                date_val = r["scan_date"]
            ws.cell(row=row_num, column=COL["Scan Date"], value=date_val).number_format = "yyyy-mm-dd"
            ws.cell(row=row_num, column=COL["IP Address"], value=r["ip"])
            ws.cell(row=row_num, column=COL["Hostname"], value=r["hostname"] or None)
            ws.cell(row=row_num, column=COL["Banner"], value=r["banner"] or None)
            ws.cell(row=row_num, column=COL["SSHv1 Supported"], value=r["sshv1"])
            ws.cell(row=row_num, column=COL["Weak KEX Count"], value=r["kex"])
            ws.cell(row=row_num, column=COL["Weak MAC Count"], value=r["mac"])
            ws.cell(row=row_num, column=COL["Weak HostKeyAlgo Count"], value=r["hka"])
            ws.cell(row=row_num, column=COL["Weak Cipher Count"], value=r["cip"])
            ws.cell(row=row_num, column=COL["Password Auth Enabled"], value=r["pwauth"])
            wk_col = get_column_letter(COL["Weak KEX Count"])
            cip_col = get_column_letter(COL["Weak Cipher Count"])
            sshv1_col = get_column_letter(COL["SSHv1 Supported"])
            total_col = get_column_letter(COL["Weak Algo Total"])
            ws.cell(row=row_num, column=COL["Weak Algo Total"],
                    value=f"=SUM({wk_col}{row_num}:{cip_col}{row_num})")
            ws.cell(row=row_num, column=COL["Weakness Score"],
                    value=f"=IF({sshv1_col}{row_num}=TRUE,100,0)+{total_col}{row_num}")
            ws.cell(row=row_num, column=COL["Host Key Info"], value=r["hostkey_info"] or None)

        last_row = 2 + len(rows) if rows else 2
        ws.freeze_panes = "A3"
        if rows:
            ws.auto_filter.ref = f"A2:M{last_row}"

        widths = {"A": 12, "B": 16, "C": 24, "D": 32, "E": 14, "F": 12,
                  "G": 12, "H": 18, "I": 14, "J": 16, "K": 14, "L": 14, "M": 40}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        return ws

    def add_tile(ws, top_row, left_col, label, number_formula, subtext_formula, number_format=None):
        c1 = get_column_letter(left_col)
        c2 = get_column_letter(left_col + 1)

        num_cell = ws[f"{c1}{top_row}"]
        ws.merge_cells(f"{c1}{top_row}:{c2}{top_row}")
        num_cell.value = number_formula
        num_cell.font = TILE_NUM_FONT
        num_cell.alignment = Alignment(horizontal="center")
        if number_format:
            num_cell.number_format = number_format

        lbl_cell = ws[f"{c1}{top_row + 1}"]
        ws.merge_cells(f"{c1}{top_row + 1}:{c2}{top_row + 1}")
        lbl_cell.value = label
        lbl_cell.font = LABEL_FONT
        lbl_cell.alignment = Alignment(horizontal="center")

        sub_cell = ws[f"{c1}{top_row + 2}"]
        ws.merge_cells(f"{c1}{top_row + 2}:{c2}{top_row + 2}")
        sub_cell.value = subtext_formula
        sub_cell.font = SUBTEXT_FONT
        sub_cell.alignment = Alignment(horizontal="center")

        return num_cell.coordinate

    def build_overview_sheet(wb):
        ws = wb.create_sheet("Overview")
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:I1")
        ws["A1"] = "SSH SECURITY POSTURE — SUBNET AUDIT"
        ws["A1"].font = HEADER_FONT
        ws["A1"].fill = HEADER_FILL
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        ws.merge_cells("A2:I2")
        ws["A2"] = ('="Report generated: "&TEXT(MAX(Report!$A$3:$A$100000),"mmmm d, yyyy")'
                     '&"   |   Full per-host detail on the Report tab"')
        ws["A2"].font = SUBTEXT_FONT
        ws["A2"].alignment = Alignment(horizontal="center")
        ws.row_dimensions[2].height = 18

        def pct_of_hosts(cell_ref):
            return f'=IFERROR(TEXT({cell_ref}/$B$4,"0%")&" of scanned hosts","—")'

        def coord(top_row, left_col):
            return f"{get_column_letter(left_col)}{top_row}"

        hosts_cell = add_tile(ws, 4, 2, "HOSTS SCANNED",
                               "=COUNTA(Report!$B$3:$B$100000)",
                               '="Subnet scan — "&TEXT(MAX(Report!$A$3:$A$100000),"mmm d, yyyy")')
        ws[hosts_cell].fill = NEUTRAL_FILL

        add_tile(ws, 4, 5, "HOSTS WITH SSHv1 ENABLED",
                 "=COUNTIF(Report!$E$3:$E$100000,TRUE)", pct_of_hosts(coord(4, 5)))
        add_tile(ws, 4, 8, "HOSTS WITH WEAK CIPHERS",
                 '=COUNTIF(Report!$I$3:$I$100000,">0")', pct_of_hosts(coord(4, 8)))
        add_tile(ws, 8, 2, "HOSTS WITH WEAK MACs",
                 '=COUNTIF(Report!$G$3:$G$100000,">0")', pct_of_hosts(coord(8, 2)))
        add_tile(ws, 8, 5, "HOSTS WITH PASSWORD AUTH ENABLED",
                 "=COUNTIF(Report!$J$3:$J$100000,TRUE)", pct_of_hosts(coord(8, 5)))
        add_tile(ws, 8, 8, "HOSTS WITH WEAK KEY EXCHANGE",
                 '=COUNTIF(Report!$F$3:$F$100000,">0")', pct_of_hosts(coord(8, 8)))
        add_tile(ws, 12, 2, "HOSTS WITH WEAK HOST-KEY ALGOS",
                 '=COUNTIF(Report!$H$3:$H$100000,">0")', pct_of_hosts(coord(12, 2)))

        clean_formula = ('=COUNTIFS(Report!$E$3:$E$100000,FALSE,Report!$F$3:$F$100000,0,'
                          'Report!$G$3:$G$100000,0,Report!$H$3:$H$100000,0,'
                          'Report!$I$3:$I$100000,0,Report!$J$3:$J$100000,FALSE)')
        add_tile(ws, 12, 5, "FULLY HARDENED HOSTS (CLEAN)", clean_formula, pct_of_hosts(coord(12, 5)))

        add_tile(ws, 12, 8, "HOSTS WITH ANY SSH WEAKNESS (AT RISK)",
                 "=IFERROR(($B$4-$E$12)/$B$4,0)", '=($B$4-$E$12)&" of "&$B$4&" hosts"',
                 number_format="0%")

        for coord_ in ("E4", "H4", "E8"):
            ws.conditional_formatting.add(
                coord_, CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL, font=RED_FONT))
            ws.conditional_formatting.add(
                coord_, CellIsRule(operator="equal", formula=["0"], fill=GREEN_FILL, font=GREEN_FONT))

        for coord_ in ("B8", "H8", "B12"):
            ws.conditional_formatting.add(
                coord_, FormulaRule(formula=[f"AND($B$4>0,{coord_}/$B$4>=0.5)"], fill=RED_FILL, font=RED_FONT, stopIfTrue=True))
            ws.conditional_formatting.add(
                coord_, FormulaRule(formula=[f"AND($B$4>0,{coord_}>0,{coord_}/$B$4<0.5)"], fill=AMBER_FILL, font=AMBER_FONT, stopIfTrue=True))
            ws.conditional_formatting.add(
                coord_, FormulaRule(formula=[f"{coord_}=0"], fill=GREEN_FILL, font=GREEN_FONT, stopIfTrue=True))

        ws.conditional_formatting.add(
            "E12", FormulaRule(formula=["IFERROR(E12/$B$4>=0.75,FALSE)"], fill=GREEN_FILL, font=GREEN_FONT, stopIfTrue=True))
        ws.conditional_formatting.add(
            "E12", FormulaRule(formula=["IFERROR(AND(E12/$B$4>=0.4,E12/$B$4<0.75),FALSE)"], fill=AMBER_FILL, font=AMBER_FONT, stopIfTrue=True))
        ws.conditional_formatting.add(
            "E12", FormulaRule(formula=["IFERROR(E12/$B$4<0.4,TRUE)"], fill=RED_FILL, font=RED_FONT, stopIfTrue=True))

        ws.conditional_formatting.add(
            "H12", CellIsRule(operator="greaterThanOrEqual", formula=["0.5"], fill=RED_FILL, font=RED_FONT))
        ws.conditional_formatting.add(
            "H12", FormulaRule(formula=["AND(H12>=0.25,H12<0.5)"], fill=AMBER_FILL, font=AMBER_FONT))
        ws.conditional_formatting.add(
            "H12", CellIsRule(operator="lessThan", formula=["0.25"], fill=GREEN_FILL, font=GREEN_FONT))

        for col, w in {"A": 3, "B": 16, "C": 16, "D": 3, "E": 16, "F": 16,
                        "G": 3, "H": 16, "I": 16}.items():
            ws.column_dimensions[col].width = w

        ws.merge_cells("A16:I16")
        ws["A16"] = "FINDINGS BREAKDOWN BY CATEGORY"
        ws["A16"].font = SECTION_FONT
        ws["A16"].border = THIN_BORDER

        breakdown = [
            ("SSHv1 Enabled", "=$E$4", RED_BAR),
            ("Weak Ciphers", "=$H$4", RED_BAR),
            ("Password Auth Enabled", "=$E$8", RED_BAR),
            ("Weak MACs", "=$B$8", AMBER_BAR),
            ("Weak Key Exchange", "=$H$8", AMBER_BAR),
            ("Weak Host-Key Algorithms", "=$B$12", AMBER_BAR),
        ]
        for i, (label, formula, _) in enumerate(breakdown):
            r = 18 + i
            ws.cell(row=r, column=2, value=label)
            ws.cell(row=r, column=3, value=formula)

        chart = BarChart()
        chart.type = "bar"
        chart.x_axis.title = "Hosts Affected"
        chart.y_axis.majorGridlines = None
        chart.legend = None
        cats = Reference(ws, min_col=2, min_row=18, max_row=23)
        vals = Reference(ws, min_col=3, min_row=18, max_row=23)
        chart.add_data(vals, titles_from_data=False)
        chart.set_categories(cats)
        chart.series[0].data_points = [
            DataPoint(idx=i, spPr=GraphicalProperties(solidFill=color))
            for i, (_, _, color) in enumerate(breakdown)
        ]
        ws.add_chart(chart, "E18")

        ws.merge_cells("A35:I35")
        ws["A35"] = "TOP OFFENDERS — HIGHEST-RISK HOSTS (WORK THESE FIRST)"
        ws["A35"].font = SECTION_FONT
        ws["A35"].fill = PatternFill("solid", fgColor="D9D9D9")

        ws.merge_cells("A36:I36")
        ws["A36"] = ("Ranked by Weakness Score = (100 if SSHv1 supported) + weak KEX + weak MAC "
                     "+ weak host-key + weak cipher counts. Password auth shown for context, not "
                     "scored. Static snapshot from this scan run (precomputed, not a live formula).")
        ws["A36"].font = SUBTEXT_FONT

        headers = ["Rank", "IP Address", "Hostname", "Scan Date", "SSHv1",
                   "Weak Algo Total", "Password Auth", "Weakness Score"]
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=37, column=c, value=h)
            cell.font = Font(bold=True)
            cell.border = THIN_BORDER

        top10 = sorted(rows, key=lambda r: r["score"], reverse=True)[:10]
        for i, r in enumerate(top10):
            row_num = 38 + i
            ws.cell(row=row_num, column=1, value=i + 1)
            ws.cell(row=row_num, column=2, value=r["ip"])
            ws.cell(row=row_num, column=3, value=r["hostname"] or "(no reverse DNS)")
            date_val = r["scan_date"]
            try:
                date_val = datetime.strptime(r["scan_date"], "%Y-%m-%d").date()
            except (ValueError, TypeError):
                pass
            c4 = ws.cell(row=row_num, column=4, value=date_val)
            c4.number_format = "yyyy-mm-dd"
            ws.cell(row=row_num, column=5, value=r["sshv1"])
            ws.cell(row=row_num, column=6, value=r["weak_total"])
            ws.cell(row=row_num, column=7, value=r["pwauth"])
            ws.cell(row=row_num, column=8, value=r["score"])

        if top10:
            last = 37 + len(top10)
            ws.conditional_formatting.add(
                f"E38:E{last}", CellIsRule(operator="equal", formula=["TRUE"], fill=RED_FILL, font=Font(color="FFFFFF", bold=True)))
            ws.conditional_formatting.add(
                f"H38:H{last}", ColorScaleRule(
                    start_type="min", start_color="C6EFCE",
                    mid_type="percentile", mid_value=50, mid_color="FFEB9C",
                    end_type="max", end_color="FFC7CE"))
            footnote_row = last + 2
        else:
            footnote_row = 39

        ws.cell(row=footnote_row, column=1,
                value="Full per-host, per-algorithm detail (including hosts beyond the top 10): see Report sheet.")
        ws.cell(row=footnote_row, column=1).font = SUBTEXT_FONT
        return ws

    wb = Workbook()
    wb.remove(wb.active)
    build_overview_sheet(wb)
    build_report_sheet(wb)
    wb.active = 0
    return wb


def write_xlsx(rows, xlsx_path):
    try:
        wb = build_workbook(rows)
    except ImportError:
        print("[!] openpyxl not installed - skipping .xlsx generation. "
              "Install with: pip install openpyxl", file=sys.stderr)
        return
    wb.save(xlsx_path)
    print(f"[*] Workbook saved to: {xlsx_path} ({len(rows)} host(s))")


# ===========================================================================
# main
# ===========================================================================

def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("subnet", nargs="?", default="192.168.0.1/24")
    ap.add_argument("--keep-temp", action="store_true", help="keep the raw nmap XML output")
    ap.add_argument("--no-xlsx", action="store_true", help="skip .xlsx generation, CSV only")
    ap.add_argument("--csv-out", default="insecure_ssh_hosts.csv")
    ap.add_argument("--xlsx-out", default="insecure_ssh_hosts.xlsx")
    ap.add_argument("--from-csv", metavar="FILE",
                     help="skip scanning; rebuild the .xlsx from an existing CSV instead")
    args = ap.parse_args()

    if args.from_csv:
        rows = read_rows_from_csv(args.from_csv)
        write_xlsx(rows, args.xlsx_out)
        return

    scan_date = datetime.now().strftime("%Y-%m-%d")
    xml_path = "nmap_temp_output.xml"

    run_nmap(args.subnet, xml_path)
    rows = parse_xml_file(xml_path, scan_date)
    print(f"[*] Hosts with SSH open: {len(rows)}")

    if args.keep_temp:
        print(f"[*] Raw nmap XML kept at: {xml_path}")
    else:
        try:
            os.remove(xml_path)
        except OSError:
            pass

    write_csv(rows, args.csv_out)

    if not args.no_xlsx:
        write_xlsx(rows, args.xlsx_out)


if __name__ == "__main__":
    main()
